"""dlt resource for the legacy Excel tracker.

Each weekly workbook has one sheet per team, in two layouts that were set up
by different people (different column names, header in row 1 or row 2 under
a title). This resource finds the header row, maps both layouts to one set of
columns, skips blank rows and loads everything as raw text. Typing, status
mapping, date parsing and de-duplication happen later in dbt, where they can
be tested.

Rows are keyed by (snapshot_week, sheet_name, source_row), so re-running on
the same files is idempotent. Each weekly file is written once and never
changes, so the resource remembers the last week it loaded (in dlt's
pipeline state) and only reads newer workbooks.
"""

from __future__ import annotations

import datetime as dt
import re
from collections.abc import Iterator
from pathlib import Path

import dlt
from openpyxl import load_workbook

# canonical column -> header spellings seen in the tracker (compared case-insensitively)
HEADER_ALIASES: dict[str, set[str]] = {
    "row_no": {"no.", "#"},
    "project_ref": {"project code", "ref"},
    "project_name": {"project name", "project"},
    "owner": {"pic", "owner"},
    "status": {"status", "stage"},
    "start_raw": {"start", "start date"},
    "end_raw": {"end", "due date"},
    "extensions": {"extensions", "ext."},
    "latest_update": {"latest update", "comments"},
}
ALIAS_TO_COLUMN = {alias: col for col, aliases in HEADER_ALIASES.items() for alias in aliases}
FILE_PATTERN = re.compile(r"PMO_tracker_(\d{4}-W\d{2})\.xlsx$")
MAX_HEADER_SEARCH_ROWS = 5


def _text(value) -> str | None:
    """Raw cell value as text; real Excel dates become ISO dates."""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """Index of the header row and a map canonical column -> column index."""
    for i, row in enumerate(rows[:MAX_HEADER_SEARCH_ROWS]):
        mapping = {}
        for j, cell in enumerate(row):
            col = ALIAS_TO_COLUMN.get((_text(cell) or "").lower())
            if col and col not in mapping:
                mapping[col] = j
        if {"project_ref", "project_name", "status"} <= mapping.keys():
            return i, mapping
    return None


def parse_sheet(rows: list[tuple]) -> tuple[str, list[dict]]:
    """Return (layout, records) for one sheet's rows."""
    found = find_header(rows)
    if not found:
        return "unknown", []
    header_idx, mapping = found
    layout = "front" if (_text(rows[header_idx][mapping["project_ref"]]) or "").lower() == "project code" else "middle"
    records = []
    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):  # 1-based Excel row number
        values = {col: _text(row[j]) if j < len(row) else None for col, j in mapping.items()}
        if not any(values.values()):
            continue
        records.append({"source_row": offset, **values})
    return layout, records


@dlt.resource(name="tracker_rows", write_disposition="merge", primary_key=("snapshot_week", "sheet_name", "source_row"))
def tracker_rows(folder: str | Path) -> Iterator[dict]:
    state = dlt.current.resource_state()
    last_week = state.get("last_week", "0000-W00")
    for path in sorted(Path(folder).glob("PMO_tracker_*.xlsx")):
        m = FILE_PATTERN.search(path.name)
        if not m or m.group(1) <= last_week:
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                layout, records = parse_sheet(list(ws.iter_rows(values_only=True)))
                for rec in records:
                    yield {"snapshot_week": m.group(1), "file_name": path.name, "sheet_name": ws.title, "layout": layout, **rec}
        finally:
            wb.close()
        state["last_week"] = m.group(1)
