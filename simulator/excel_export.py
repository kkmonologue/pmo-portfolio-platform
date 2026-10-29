"""Weekly export of the legacy Excel project tracker.

Before the PMO system went live, the portfolio lived in one Excel workbook
with a sheet per team. Teams kept updating it in parallel afterwards, so it
is a second, partly stale source that has to be reconciled with the database:

* front-office (regional) sheets and middle-office (HQ) sheets were set up by
  different people, so they use different column names, header positions and
  date formats (text "dd/mm/yyyy", real Excel dates, or "30.06.2025" strings);
* statuses use the old tracker vocabulary ("In Progress", "Close report
  pending", "Backlog", ...), sometimes blank;
* about one row in five lags two to five weeks behind the system;
* projects started before go-live are listed under their legacy TRK code;
* stray blank rows and duplicated rows appear, as in any shared workbook.
"""

from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .company import TEAMS

LEGACY_STATUS = {
    "not_started": ["Backlog", "Pending", "Not started"],
    "active": ["In Progress", "In Progress", "Ongoing"],
    "on_hold": ["Pending", "On hold"],
    "to_report": ["Close report pending", "To report"],
    "closed": ["Closed", "Closed", "Done"],
    "canceled": ["Canceled", "Cancelled"],
}
SPEC = {t["code"]: t for t in TEAMS}
FRONT_HEAD = ["No.", "Project Code", "Project Name", "PIC", "Status", "Start", "End", "Extensions", "Latest update"]
MIDDLE_HEAD = ["#", "Ref", "Project", "Owner", "Stage", "Start date", "Due date", "Ext.", "Comments"]


def _middle_date(rng: random.Random, d: dt.date | None):
    if not d:
        return ""
    return d if rng.random() < 0.7 else d.strftime("%d.%m.%Y")


def write_week(out_dir: Path, monday: dt.date, snapshot: dict, seed: int) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(f"{seed}|tracker-xl|{monday}")
    wb = Workbook()
    wb.remove(wb.active)
    for code in [t["code"] for t in TEAMS]:
        rows = snapshot.get(code, [])
        front = SPEC[code]["group"] == "R"
        ws = wb.create_sheet(SPEC[code]["en"][:31] if front else f"HQ - {SPEC[code]['en']}"[:31])
        if not front and rng.random() < 0.5:
            ws.append([f"{SPEC[code]['en']} - project tracker"])
            ws["A1"].font = Font(bold=True, size=13)
        ws.append(FRONT_HEAD if front else MIDDLE_HEAD)
        for i, r in enumerate(rows, 1):
            status = "" if rng.random() < 0.04 else rng.choice(LEGACY_STATUS[r["status"]])
            if front:
                start = r["start"].strftime("%d/%m/%Y") if r["start"] else "TBD"
                end = r["end"].strftime("%d/%m/%Y") if r["end"] else "TBD"
            else:
                start, end = _middle_date(rng, r["start"]), _middle_date(rng, r["end"])
            update = " ".join(r["update"].split())[:300]
            ws.append([i, r["legacy_code"] or r["code"], r["title"], r["owner"], status, start, end, r["ext"] or None, update])
            for col in (6, 7):
                if isinstance(ws.cell(ws.max_row, col).value, dt.date):
                    ws.cell(ws.max_row, col).number_format = "yyyy-mm-dd"
        if rows and rng.random() < 0.3:
            ws.insert_rows(rng.randint(3, ws.max_row + 1), 1)
        if rows and rng.random() < 0.15:
            ws.append([c.value for c in ws[ws.max_row]])
    iso = monday.isocalendar()
    path = out_dir / f"PMO_tracker_{iso.year}-W{iso.week:02d}.xlsx"
    wb.save(path)
    return [path]
