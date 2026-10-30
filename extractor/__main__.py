"""Command line.

  python -m extractor eval                                  # rules on both evaluation sets
  python -m extractor eval --extractor rules llm --provider gemini --sample 300
  python -m extractor extract --input updates.xlsx --text-column Comments --date-column Date --output signals.csv
  python -m extractor warehouse --extractor rules           # incremental, into the DuckDB warehouse

API keys come from environment variables or a .env file in the current folder
(GEMINI_API_KEY, GROQ_API_KEY, MISTRAL_API_KEY or LLM_API_KEY). Never commit them.
"""

import argparse
import csv
import datetime as dt
import json
import os
import sys
from functools import partial
from pathlib import Path

from dotenv import load_dotenv

from . import rules
from .categories import load_categories
from .evaluate import errors, load_challenge, load_simulated, markdown_report, run, score, stratified_sample
from .schema import Extraction

DATA_DIR = Path(os.getenv("DATA_DIR", "data"))


def make_extractor(kind: str, provider: str, model: str | None, categories: dict, rpm: int | None = None):
    """(name, callable) for 'rules' or 'llm'."""
    if kind == "rules":
        return "rules", partial(rules.extract, categories=categories)
    from .llm import LLMExtractor

    llm = LLMExtractor(provider=provider, model=model, categories=categories, rpm=rpm, cache_path=DATA_DIR / "llm_cache.sqlite")
    return llm.name, llm.extract


def _progress(i: int, n: int) -> None:
    if i == n or i % 25 == 0:
        print(f"  {i}/{n}", file=sys.stderr)


def cmd_eval(args) -> None:
    categories = load_categories(args.categories)
    sets = {"challenge": load_challenge()}
    labels = DATA_DIR / "labels" / "status_updates.jsonl"
    source = os.getenv("SOURCE_DB_URL", f"sqlite:///{DATA_DIR / 'source.db'}")
    if labels.exists():
        sets["simulated"] = stratified_sample(load_simulated(source, labels), args.sample, args.seed)
    else:
        print("No simulator labels found; run `python -m simulator init` to evaluate on simulated updates too.", file=sys.stderr)
    results, details = {}, {}
    for kind in args.extractor:
        name, fn = make_extractor(kind, args.provider, args.model, categories, args.rpm)
        for set_name, items in sets.items():
            print(f"{name} on {set_name} ({len(items)} updates)", file=sys.stderr)
            preds = run(items, fn, _progress)
            results.setdefault(set_name, {})[name] = score(items, preds)
            details.setdefault(set_name, {})[name] = errors(items, preds)
    out = DATA_DIR / "eval"
    out.mkdir(parents=True, exist_ok=True)
    meta = {"generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "sample": args.sample, "seed": args.seed,
            "note": f"Simulated set: stratified sample of {args.sample} labelled updates (seed {args.seed}). Challenge set: {len(sets['challenge'])} hand-labelled updates."}  # fmt: skip
    (out / "report.md").write_text(markdown_report(results, meta), encoding="utf-8")
    (out / "results.json").write_text(
        json.dumps({"meta": meta, "results": results, "errors": details}, indent=2, default=str), encoding="utf-8"
    )
    print((out / "report.md").read_text(encoding="utf-8"))
    print(f"\nwritten: {out / 'report.md'}, {out / 'results.json'}", file=sys.stderr)


def _read_table(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else f"column_{i}" for i, h in enumerate(rows[0])]
        return [dict(zip(header, r, strict=False)) for r in rows[1:] if any(v is not None for v in r)]
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def _write_table(path: Path, rows: list[dict]) -> None:
    cols = list(rows[0]) if rows else []
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        wb.save(path)
    else:
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)


def _to_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def cmd_extract(args) -> None:
    categories = load_categories(args.categories)
    name, fn = make_extractor(args.extractor, args.provider, args.model, categories, args.rpm)
    rows = _read_table(Path(args.input))
    if rows and args.text_column not in rows[0]:
        raise SystemExit(f"Column {args.text_column!r} not found; columns are: {', '.join(rows[0])}")
    today = dt.date.today()
    for i, row in enumerate(rows, 1):
        text = str(row.get(args.text_column) or "")
        posted = (_to_date(row.get(args.date_column)) if args.date_column else None) or today
        e = fn(text, posted) if text.strip() else Extraction()
        row.update({"rag": e.rag, "stated_rag": e.stated_rag, "blocked": e.blocked,
                    "risks": "; ".join(f"{r.category} ({r.severity})" for r in e.risks),
                    "next_milestone_date": e.next_milestone_date, "date_precision": e.next_milestone_date_precision})  # fmt: skip
        _progress(i, len(rows))
    _write_table(Path(args.output), rows)
    print(f"{len(rows)} rows with {name} signals written to {args.output}")


def cmd_warehouse(args) -> None:
    from .warehouse import extract_new

    name, fn = make_extractor(args.extractor, args.provider, args.model, load_categories(args.categories), args.rpm)
    counts = extract_new(Path(args.warehouse), name, fn, args.max_new)
    print(f"{name}: extracted {counts['extracted']} new updates, {counts['remaining']} still to do")


def main() -> None:
    load_dotenv()
    ap = argparse.ArgumentParser(prog="extractor", description="Structured signals from free-text project status updates")
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--provider", default=os.getenv("LLM_PROVIDER", "gemini"), help="gemini | groq | mistral | ollama")
    common.add_argument("--model", default=None, help="override the provider's default model")
    common.add_argument("--rpm", type=int, default=None, help="max requests per minute")
    common.add_argument("--categories", default=None, help="JSON file with your own risk categories")
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("eval", parents=[common], help="score extractors on the evaluation sets")
    e.add_argument("--extractor", nargs="+", default=["rules"], choices=["rules", "llm"])
    e.add_argument("--sample", type=int, default=300)
    e.add_argument("--seed", type=int, default=7)
    x = sub.add_parser("extract", parents=[common], help="add signals to any CSV / Excel file of updates")
    x.add_argument("--input", required=True)
    x.add_argument("--text-column", required=True)
    x.add_argument("--date-column", default=None, help="when each update was written (default: today)")
    x.add_argument("--extractor", default="rules", choices=["rules", "llm"])
    x.add_argument("--output", required=True)
    w = sub.add_parser("warehouse", parents=[common], help="extract new updates into the DuckDB warehouse")
    w.add_argument("--extractor", default="rules", choices=["rules", "llm"])
    w.add_argument("--warehouse", default=str(DATA_DIR / "warehouse.duckdb"))
    w.add_argument("--max-new", type=int, default=None, help="cap per run (useful on free API tiers)")
    args = ap.parse_args()
    {"eval": cmd_eval, "extract": cmd_extract, "warehouse": cmd_warehouse}[args.cmd](args)


if __name__ == "__main__":
    main()
